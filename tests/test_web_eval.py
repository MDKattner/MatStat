"""Tests for the web Team Evaluation flow — the job body and API routes."""

import sys
import time
from pathlib import Path
from typing import Any

import openpyxl
import pytest

sys.path.append(str(Path(__file__).parent.parent))

from scripts.web import app as web_app
from scripts.web import team_eval

CSV_ROW: str = '"vid.mkv:1",0,6,A,collar tie,"double","sprawl","T","E"'
CSV_ROW2: str = '"vid.mkv:2",6,12,D,standing,"sprawl","sweep single","E","T"'


class FakeContext:
    """Minimal stand-in for scripts.web.jobs.JobContext."""

    def __init__(self) -> None:
        self.job_id: str = "testjob"
        self.reports: list[tuple[float, str]] = []

    def report(self, progress: float, message: str = "") -> None:
        self.reports.append((progress, message))

    @property
    def cancelled(self) -> bool:
        return False


def _redirect_dirs(monkeypatch, tmp_path: Path) -> dict[str, Path]:
    """Point the team_eval and web_app data dirs at a temp dir."""
    dirs: dict[str, Path] = {
        "csv": tmp_path / "csv",
        "eval": tmp_path / "eval",
    }
    for directory in dirs.values():
        directory.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(team_eval, "csv_dir", dirs["csv"])
    monkeypatch.setattr(team_eval, "eval_dir", dirs["eval"])
    monkeypatch.setattr(web_app, "csv_dir", dirs["csv"])
    monkeypatch.setattr(web_app, "eval_dir", dirs["eval"])
    return dirs


def _write_csv(tmp_path: Path, name: str, content: str = CSV_ROW) -> None:
    (tmp_path / "csv" / f"{name}.csv").write_text(content)


def _wait_for_job(client, job_id: str, timeout: float = 5.0) -> dict[str, Any]:
    deadline: float = time.monotonic() + timeout
    while time.monotonic() < deadline:
        resp = client.get(f"/api/jobs/{job_id}")
        assert resp.status_code == 200
        job: dict[str, Any] = resp.json()
        if job["status"] in ("done", "failed", "cancelled"):
            return job
        time.sleep(0.01)
    raise AssertionError(f"Job {job_id} did not finish in {timeout}s")


class TestRunTeamEvalJob:
    """Tests for team_eval.run_team_eval_job — the Excel generation job."""

    def test_writes_excel_with_sheet_per_wrestler(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "Alice")
        _write_csv(tmp_path, "Bob Smith")
        _write_csv(tmp_path, "UNKNOWN")

        ctx: FakeContext = FakeContext()
        result: dict[str, Any] = team_eval.run_team_eval_job(ctx)

        assert result == {
            "output": "Team_Stats.xlsx",
            "processed": ["Alice", "Bob Smith"],
            "errors": [],
        }
        report: Path = tmp_path / "eval" / "Team_Stats.xlsx"
        assert report.is_file()
        wb = openpyxl.load_workbook(report)
        assert wb.sheetnames == ["Team Summary", "Alice", "Bob Smith"]
        assert ctx.reports[-1] == (100, "Report generated: Team_Stats.xlsx")

    def test_section_labels_and_layout(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "Alice")

        team_eval.run_team_eval_job(FakeContext())

        ws = openpyxl.load_workbook(tmp_path / "eval" / "Team_Stats.xlsx")["Alice"]
        assert ws["A1"].value == "Initiation"
        assert ws["A1"].font.bold is True
        assert ws["A7"].value == "Rates"
        assert ws["A10"].value == "Defense"
        assert ws["I10"].value == "Offense"
        assert ws["Q10"].value == "Raw Data"
        assert ws["A2"].value == "Segment"
        assert ws["A8"].value == "Metric"
        assert ws["A11"].value == "Move"
        assert ws["I11"].value == "Move"
        assert ws["Q11"].value == "Origin"

    def test_team_summary_sheet(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "Alice")
        _write_csv(tmp_path, "Bob", content=CSV_ROW2)

        team_eval.run_team_eval_job(FakeContext())

        ws = openpyxl.load_workbook(tmp_path / "eval" / "Team_Stats.xlsx")[
            "Team Summary"
        ]
        assert ws["A1"].value == "Wrestler"
        assert ws["B1"].value == "Matches"
        assert ws["C1"].value == "Net Points per Match"
        assert ws["I1"].value == "Net Points per Match (z)"
        assert ws["A2"].value == "Alice"
        assert ws["A3"].value == "Bob"
        assert ws["A4"].value == "Team Mean"

    def test_initiation_segments_in_report(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        content: str = (
            '"a.mkv:1",0,6,A,collar tie,"double","sprawl","T","E",2,2,W\n'
            '"a.mkv:2",6,12,D,standing,"sprawl","sweep single","E","T",-4,-4,L\n'
            '"a.mkv:3",12,18,A,front headlock,"single","whizzer","T","",3,3,\n'
        )
        _write_csv(tmp_path, "Alice", content=content)

        team_eval.run_team_eval_job(FakeContext())

        ws = openpyxl.load_workbook(tmp_path / "eval" / "Team_Stats.xlsx")["Alice"]
        assert ws["A2"].value == "Segment"
        assert ws["A3"].value == "All"
        assert ws["A4"].value == "Wins"
        assert ws["A5"].value == "Losses"
        assert ws["A6"].value == "Unrecorded"
        # W row is attacking (count 1); L row is not. "Sequences" is column B,
        # so "Attack Count" is column C.
        assert ws["B4"].value == 1
        assert ws["B5"].value == 1
        assert ws["B6"].value == 1
        assert ws["C4"].value == 1
        assert ws["C5"].value == 0
        # All includes all three rows, two of which attack.
        assert ws["B3"].value == 3
        assert ws["C3"].value == 2
        # The unrecorded row is attacking (count 1).
        assert ws["C6"].value == 1

    def test_per_file_error_keeps_others(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "Alice")
        _write_csv(tmp_path, "Broken", content='"v.mkv:1",abc,6,A,collar tie,"double","sprawl","T","E"')

        result: dict[str, Any] = team_eval.run_team_eval_job(FakeContext())

        assert result["processed"] == ["Alice"]
        assert len(result["errors"]) == 1
        assert "Broken" in result["errors"][0]
        wb = openpyxl.load_workbook(tmp_path / "eval" / "Team_Stats.xlsx")
        assert wb.sheetnames == ["Team Summary", "Alice"]

    def test_no_csv_files_raises(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)

        with pytest.raises(ValueError, match="No wrestler data files"):
            team_eval.run_team_eval_job(FakeContext())

    def test_skips_unknown_csv(self, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "UNKNOWN")

        with pytest.raises(ValueError, match="No wrestler data files"):
            team_eval.run_team_eval_job(FakeContext())


class TestTeamEvalRoute:
    """Tests for POST /api/team-eval and GET /api/reports/{file}."""

    def test_job_completes_and_download_works(self, client, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)
        _write_csv(tmp_path, "Alice")

        resp = client.post("/api/team-eval")
        assert resp.status_code == 200
        body: dict[str, Any] = resp.json()
        assert body["status"] == "queued"
        job_id: str = body["job_id"]

        job: dict[str, Any] = _wait_for_job(client, job_id)
        assert job["status"] == "done"
        assert job["result"]["processed"] == ["Alice"]
        assert (tmp_path / "eval" / "Team_Stats.xlsx").is_file()

        dl = client.get("/api/reports/Team_Stats.xlsx")
        assert dl.status_code == 200
        assert dl.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert dl.content.startswith(b"PK")  # xlsx is a zip archive

    def test_no_data_404(self, client, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)

        resp = client.post("/api/team-eval")
        assert resp.status_code == 404
        assert "No wrestler data" in resp.json()["detail"]

    def test_missing_report_404(self, client, tmp_path, monkeypatch) -> None:
        _redirect_dirs(monkeypatch, tmp_path)

        resp = client.get("/api/reports/Team_Stats.xlsx")
        assert resp.status_code == 404

    def test_invalid_report_name_400(self, client) -> None:
        resp = client.get("/api/reports/.hidden.xlsx")
        assert resp.status_code == 400
